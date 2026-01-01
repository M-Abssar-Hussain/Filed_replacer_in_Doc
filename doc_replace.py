from docx import Document

def replace_placeholders(doc, replacements):
    for paragraph in doc.paragraphs:
        for key, value in replacements.items():
            if key in paragraph.text:
                paragraph.text = paragraph.text.replace(key, value)


def main():
    name = input("Enter name: ")
    phone = input("Enter phone number: ")
    email = input("Enter email: ")

    replacements = {
        "{name}": name,
        "{phone}": phone,
        "{email}": email
    }


    input_file = r"Enter file path"
    output_file = r"Enter file path"

    doc = Document(input_file)
    replace_placeholders(doc, replacements)
    doc.save(output_file)

    print("Document updated successfully!")


if __name__ == "__main__":
    main()
